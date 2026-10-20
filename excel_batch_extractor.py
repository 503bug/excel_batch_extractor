# -*- coding: utf-8 -*-
"""
批量提取多个 Excel 基础数据工作簿中
指定工作表、指定单元格的数据（可多项），
写入汇总表并标注数据来源（工作簿名称）。

GUI 窗口版（单窗口操作，带进度条与日志）

依赖：pip install openpyxl
"""

import os
import re
import threading
import warnings
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# 屏蔽 openpyxl 对“未知扩展格式”的 UserWarning（不影响数据提取）
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

from openpyxl import load_workbook, Workbook
from openpyxl.utils import coordinate_to_tuple, get_column_letter

CELL_PATTERN = re.compile(r"^[A-Za-z]{1,3}[1-9]\d*$")


# ---------------- 单元格读取 ----------------
def read_cell_value(ws, coord):
    """读取单元格值；若为合并单元格，取左上角单元格的值。"""
    coord = coord.upper()
    row, col = coordinate_to_tuple(coord)
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            coord = rng.start_cell.coordinate
            break
    return ws[coord].value


# ---------------- 解析提取项 ----------------
def parse_items(raw):
    """
    每行一个提取项，格式：工作表名!单元格
    支持全角感叹号，返回 (items, err)：
      items = [(工作表名, 单元格), ...]
    """
    items = []
    for line in raw.splitlines():
        line = line.strip()
        if not line:
            continue
        line = line.replace("！", "!")
        if "!" not in line:
            return None, f"格式错误（缺少“!”）：{line}"
        sheet, cell = line.rsplit("!", 1)
        sheet = sheet.strip()
        cell = cell.strip().upper()
        if not sheet:
            return None, f"工作表名为空：{line}"
        if not CELL_PATTERN.match(cell):
            return None, f"单元格格式错误：{line}"
        items.append((sheet, cell))
    if not items:
        return None, "未输入任何提取项"
    return items, None


# ---------------- 数据提取 ----------------
def collect_data(src_dir, items, progress=None):
    """遍历文件夹下所有 xlsx/xlsm，按 items 提取数据。"""
    files = [
        f for f in sorted(os.listdir(src_dir))
        if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$")
    ]

    results = []
    total = len(files)
    for idx, fname in enumerate(files, 1):
        if progress:
            progress(idx, total, fname)

        base_name = os.path.splitext(fname)[0]
        fpath = os.path.join(src_dir, fname)

        values, notes = [], []
        try:
            wb_val = load_workbook(fpath, data_only=True)
            wb_raw = None  # 延迟加载，仅在需要判断公式时使用
            try:
                for sheet_name, cell_ref in items:
                    value, note = None, ""
                    if sheet_name not in wb_val.sheetnames:
                        note = "工作表不存在"
                    else:
                        value = read_cell_value(wb_val[sheet_name], cell_ref)
                        if value is None:
                            if wb_raw is None:
                                wb_raw = load_workbook(fpath, data_only=False)
                            if sheet_name in wb_raw.sheetnames:
                                raw = read_cell_value(wb_raw[sheet_name], cell_ref)
                                if isinstance(raw, str) and raw.startswith("="):
                                    note = "公式无缓存结果"
                    values.append(value)
                    notes.append(note)
            finally:
                wb_val.close()
                if wb_raw is not None:
                    wb_raw.close()
        except Exception as e:
            values = [None] * len(items)
            notes = [f"读取失败：{e}"] * len(items)

        results.append((base_name, values, notes))
    return results


# ---------------- 写入汇总表 ----------------
def write_results(results, items, out_path, append_mode):
    """将结果写入汇总表，返回 (ok_count, problems)。"""
    name_col = 1        # A 列
    value_col = 2       # B 列

    if append_mode and os.path.exists(out_path):
        out_wb = load_workbook(out_path)
        out_ws = out_wb.active
        first_row_empty = all(c.value is None for c in out_ws[1])
        if first_row_empty:
            write_row, header_row = 2, 1
        else:
            write_row, header_row = out_ws.max_row + 1, None
    else:
        out_wb = Workbook()
        out_ws = out_wb.active
        out_ws.title = "汇总"
        write_row, header_row = 2, 1

    # 表头
    if header_row is not None:
        out_ws.cell(row=header_row, column=name_col, value="工作簿名称")
        for i, (sheet, cell) in enumerate(items):
            out_ws.cell(row=header_row, column=value_col + i,
                        value=f"{sheet}!{cell}")

    # 数据
    problems = []
    ok_count = 0
    for name, values, notes in results:
        out_ws.cell(row=write_row, column=name_col, value=name)
        row_ok = True
        for i, (v, nt) in enumerate(zip(values, notes)):
            out_ws.cell(row=write_row, column=value_col + i, value=v)
            if nt:
                row_ok = False
                s, c = items[i]
                problems.append(f"{name} → {s}!{c}：{nt}")
        if row_ok:
            ok_count += 1
        write_row += 1

    # 列宽
    out_ws.column_dimensions[get_column_letter(name_col)].width = 22
    for i in range(len(items)):
        letter = get_column_letter(value_col + i)
        cur = out_ws.column_dimensions[letter].width or 0
        out_ws.column_dimensions[letter].width = max(cur, 16)

    out_wb.save(out_path)
    return ok_count, problems


# ---------------- GUI ----------------
class ExcelExtractorGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel 数据批量提取工具")
        self.root.geometry("820x760")
        self.root.minsize(720, 640)

        self._running = False
        self._build_ui()

    # ---------- 界面 ----------
    def _build_ui(self):
        main = ttk.Frame(self.root, padding=12)
        main.pack(fill="both", expand=True)

        # --- 1. 基础数据文件夹 ---
        f1 = ttk.LabelFrame(
            main, text=" 1. 基础数据文件夹（待提取的 Excel 所在目录） ", padding=8)
        f1.pack(fill="x", pady=(0, 8))

        self.src_dir_var = tk.StringVar()
        ttk.Entry(f1, textvariable=self.src_dir_var).pack(
            side="left", fill="x", expand=True)
        ttk.Button(f1, text="浏览…", width=10,
                   command=self.choose_src_dir).pack(side="left", padx=(8, 0))

        # --- 2. 提取项 ---
        f2 = ttk.LabelFrame(
            main, text=" 2. 提取项（每行一个，格式：工作表名!单元格） ", padding=8)
        f2.pack(fill="both", expand=True, pady=(0, 8))

        ttk.Label(f2, text="示例：Sheet1!O7    支持全角“！”，重复项会生成多列。",
                  foreground="#666").pack(anchor="w")

        text_wrap = ttk.Frame(f2)
        text_wrap.pack(fill="both", expand=True, pady=(4, 0))

        self.items_text = tk.Text(text_wrap, height=8, wrap="none",
                                  font=("Consolas", 10))
        y_scroll = ttk.Scrollbar(text_wrap, orient="vertical",
                                 command=self.items_text.yview)
        self.items_text.configure(yscrollcommand=y_scroll.set)
        self.items_text.pack(side="left", fill="both", expand=True)
        y_scroll.pack(side="right", fill="y")
        self.items_text.insert("1.0", "Sheet1!O7")

        # --- 3. 输出设置 ---
        f3 = ttk.LabelFrame(main, text=" 3. 汇总表保存设置 ", padding=8)
        f3.pack(fill="x", pady=(0, 8))

        row1 = ttk.Frame(f3)
        row1.pack(fill="x", pady=(0, 6))
        ttk.Label(row1, text="保存文件夹：", width=12).pack(side="left")
        self.out_dir_var = tk.StringVar()
        ttk.Entry(row1, textvariable=self.out_dir_var).pack(
            side="left", fill="x", expand=True)
        ttk.Button(row1, text="浏览…", width=10,
                   command=self.choose_out_dir).pack(side="left", padx=(8, 0))

        row2 = ttk.Frame(f3)
        row2.pack(fill="x", pady=(0, 6))
        ttk.Label(row2, text="文件名：", width=12).pack(side="left")
        self.out_name_var = tk.StringVar(value="汇总表")
        ttk.Entry(row2, textvariable=self.out_name_var).pack(
            side="left", fill="x", expand=True)

        row3 = ttk.Frame(f3)
        row3.pack(fill="x")
        ttk.Label(row3, text="已存在时：", width=12).pack(side="left")
        self.exist_var = tk.StringVar(value="ask")
        ttk.Radiobutton(row3, text="询问", variable=self.exist_var,
                        value="ask").pack(side="left")
        ttk.Radiobutton(row3, text="追加到末尾", variable=self.exist_var,
                        value="append").pack(side="left", padx=(10, 0))
        ttk.Radiobutton(row3, text="覆盖重建", variable=self.exist_var,
                        value="overwrite").pack(side="left", padx=(10, 0))

        # --- 操作按钮 ---
        btn_bar = ttk.Frame(main)
        btn_bar.pack(fill="x", pady=(0, 6))

        self.start_btn = ttk.Button(btn_bar, text="开始提取",
                                    command=self.start)
        self.start_btn.pack(side="left")
        ttk.Button(btn_bar, text="清空日志",
                   command=self.clear_log).pack(side="left", padx=(8, 0))
        ttk.Button(btn_bar, text="退出",
                   command=self.root.destroy).pack(side="right")

        # --- 进度条与状态 ---
        self.progress = ttk.Progressbar(main, mode="determinate", maximum=1)
        self.progress.pack(fill="x", pady=(0, 4))

        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(main, textvariable=self.status_var,
                  foreground="#333").pack(anchor="w")

        # --- 日志 ---
        f4 = ttk.LabelFrame(main, text=" 日志 ", padding=8)
        f4.pack(fill="both", expand=True)

        self.log_text = tk.Text(f4, height=8, wrap="word",
                                font=("Consolas", 9), state="disabled")
        log_scroll = ttk.Scrollbar(f4, orient="vertical",
                                   command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=log_scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        log_scroll.pack(side="right", fill="y")

    # ---------- 路径选择 ----------
    def choose_src_dir(self):
        d = filedialog.askdirectory(title="请选择【基础数据工作簿】所在的文件夹")
        if d:
            self.src_dir_var.set(d)
            if not self.out_dir_var.get().strip():
                self.out_dir_var.set(d)

    def choose_out_dir(self):
        d = filedialog.askdirectory(title="请选择【汇总表】保存的文件夹")
        if d:
            self.out_dir_var.set(d)

    # ---------- 日志 / 状态 ----------
    def log(self, msg):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", msg + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def clear_log(self):
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")

    def _set_status(self, text):
        self.status_var.set(text)

    # ---------- 开始 ----------
    def start(self):
        if self._running:
            return

        # 1) 基础数据文件夹
        src_dir = self.src_dir_var.get().strip()
        if not src_dir:
            messagebox.showwarning("提示", "请先选择基础数据文件夹。")
            return
        if not os.path.isdir(src_dir):
            messagebox.showerror("错误", f"基础数据文件夹不存在：\n{src_dir}")
            return

        # 2) 提取项
        raw = self.items_text.get("1.0", "end").strip()
        items, err = parse_items(raw)
        if err:
            messagebox.showerror("输入错误", err)
            return

        # 3) 保存文件夹
        out_dir = self.out_dir_var.get().strip()
        if not out_dir:
            messagebox.showwarning("提示", "请先选择汇总表保存文件夹。")
            return
        if not os.path.isdir(out_dir):
            messagebox.showerror("错误", f"汇总表保存文件夹不存在：\n{out_dir}")
            return

        # 4) 文件名
        out_name = self.out_name_var.get().strip() or "汇总表"
        if not out_name.lower().endswith((".xlsx", ".xlsm")):
            out_name += ".xlsx"
        out_path = os.path.join(out_dir, out_name)

        # 5) 已存在时的处理方式
        append_mode = False
        if os.path.exists(out_path):
            choice = self.exist_var.get()
            if choice == "ask":
                ans = messagebox.askyesnocancel(
                    "汇总表已存在",
                    f"汇总表已存在：\n{out_path}\n\n"
                    "【是】：追加到已有数据末尾\n"
                    "【否】：覆盖重建\n"
                    "【取消】：退出")
                if ans is None:
                    return
                append_mode = ans
            elif choice == "append":
                append_mode = True
            else:  # overwrite
                append_mode = False

        # 启动后台任务
        self._running = True
        self.start_btn.configure(state="disabled")
        self.progress.configure(maximum=1, value=0)
        self.clear_log()
        self.log(f"基础数据文件夹：{src_dir}")
        self.log(f"汇总表路径：{out_path}")
        self.log(f"提取项：{len(items)} 个")
        self.log("=" * 60)
        self.status_var.set("正在扫描文件…")

        t = threading.Thread(
            target=self._worker,
            args=(src_dir, items, out_path, append_mode),
            daemon=True)
        t.start()

    # ---------- 后台线程 ----------
    def _worker(self, src_dir, items, out_path, append_mode):
        try:
            def progress_cb(idx, total, fname):
                self.root.after(0, self._on_progress, idx, total, fname)

            results = collect_data(src_dir, items, progress=progress_cb)

            if not results:
                self.root.after(
                    0, self._finish_error,
                    f"在以下文件夹中未找到 .xlsx / .xlsm 文件：\n{src_dir}")
                return

            self.root.after(0, self._set_status, "正在写入汇总表…")
            ok_count, problems = write_results(
                results, items, out_path, append_mode)

            self.root.after(
                0, self._finish_ok,
                results, ok_count, problems, out_path, items)

        except PermissionError:
            self.root.after(
                0, self._finish_error,
                "保存失败：汇总表可能正被 Excel 打开，请关闭后重试。")
        except Exception as e:
            self.root.after(0, self._finish_error, f"处理失败：{e}")

    # ---------- 主线程回调 ----------
    def _on_progress(self, idx, total, fname):
        self.progress.configure(maximum=total, value=idx)
        self.status_var.set(f"正在处理 {idx}/{total}：{fname}")

    def _finish_ok(self, results, ok_count, problems, out_path, items):
        self._running = False
        self.start_btn.configure(state="normal")
        self.progress.configure(value=self.progress["maximum"])
        self.status_var.set("完成")

        self.log("=" * 60)
        self.log(f"扫描工作簿：{len(results)} 个")
        self.log(f"提取项数：{len(items)} 个")
        self.log(f"全部成功的工作簿：{ok_count} 个")
        self.log(f"汇总表路径：{out_path}")
        if problems:
            self.log("")
            self.log("以下提取项存在问题（对应单元格为空）：")
            for p in problems[:50]:
                self.log("  " + p)
            if len(problems) > 50:
                self.log(f"  …… 另有 {len(problems) - 50} 条")

        msg = (f"处理完成！\n\n"
               f"扫描工作簿：{len(results)} 个\n"
               f"提取项数：{len(items)} 个\n"
               f"全部成功的工作簿：{ok_count} 个\n"
               f"汇总表路径：\n{out_path}")
        if problems:
            msg += "\n\n以下提取项存在问题（对应单元格为空）：\n" + \
                   "\n".join(problems[:20])
            if len(problems) > 20:
                msg += f"\n…… 另有 {len(problems) - 20} 条"
        messagebox.showinfo("完成", msg)

    def _finish_error(self, msg):
        self._running = False
        self.start_btn.configure(state="normal")
        self.status_var.set("出错")
        self.log("错误：" + msg)
        messagebox.showerror("错误", msg)


# ---------------- 入口 ----------------
def main():
    root = tk.Tk()
    try:
        style = ttk.Style(root)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        elif "clam" in style.theme_names():
            style.theme_use("clam")
    except Exception:
        pass

    ExcelExtractorGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
